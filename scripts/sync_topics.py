#!/usr/bin/env python3
"""Copy every topic in ebook-topics/topics-log.json into ebooks-Raw/topics-list.xlsx.

Rules:
  - A topic already in the spreadsheet (same name, ignoring case and spacing) is not added twice.
  - A topic listed in scripts/topic-holds.json is not added (the owner has to decide).
  - A topic already in ebooks-Final/used-topics.json or with a folder in ebooks-Final is not added.
  - New rows get Status "Selected", so scripts/pick_topic.py can pick them.
Prints how many rows were added. Run from the repository root.
"""
import json
import os
import re
import sys

import openpyxl

LOG = "ebook-topics/topics-log.json"
SHEET = "ebooks-Raw/topics-list.xlsx"
HOLDS = "scripts/topic-holds.json"
USED = "ebooks-Final/used-topics.json"
FINAL = "ebooks-Final"


def norm(text):
    return re.sub(r"\s+", " ", str(text or "").strip()).lower()


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", norm(text)).strip("-")


def load_json(path, default):
    if not os.path.exists(path):
        return default
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def joined(value):
    if isinstance(value, list):
        return "\n".join(str(v) for v in value)
    return "" if value is None else str(value)


def main():
    log = load_json(LOG, [])
    holds = {norm(h["topic_name"]) for h in load_json(HOLDS, [])}
    used = {norm(e.get("topic")) for e in load_json(USED, [])}
    folders = {slug(n) for n in os.listdir(FINAL) if os.path.isdir(os.path.join(FINAL, n))} if os.path.isdir(FINAL) else set()

    wb = openpyxl.load_workbook(SHEET)
    ws = wb["Topics"]
    header = [str(c.value) for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(header)}
    existing = {norm(ws.cell(row=r, column=col["Topic"]).value) for r in range(2, ws.max_row + 1)}
    next_no = max([int(ws.cell(row=r, column=col["No."]).value or 0) for r in range(2, ws.max_row + 1)] + [0]) + 1

    added = 0
    for entry in log:
        candidates = []
        if entry.get("selected_topic"):
            candidates.append(entry["selected_topic"])
        candidates += entry.get("selected_topics") or []
        for t in candidates:
            if not t or not t.get("topic_name"):
                continue
            name = t["topic_name"]
            key = norm(name)
            if key in existing or key in holds or key in used or slug(name) in folders:
                print(f"skip: {name[:60]}", file=sys.stderr)
                continue
            values = {
                "No.": next_no,
                "Date": entry.get("date", ""),
                "Status": "Selected",
                "Topic": name,
                "Category": t.get("category", ""),
                "Description": t.get("description", ""),
                "Format": t.get("format", ""),
                "Price (USD)": t.get("price_usd", ""),
                "Royalty": t.get("royalty_tier", ""),
                "Pages": t.get("page_count", ""),
                "Words": t.get("word_count", ""),
                "Target reader": t.get("target_reader", ""),
                "Chapter outline": joined(t.get("chapter_outline")),
                "Title options": joined(t.get("title_options")),
                "KDP keywords": joined(t.get("kdp_keywords")),
                "Risks": joined(t.get("risks")),
                "Sales estimate (ESTIMATE)": joined(t.get("sales_estimate")),
            }
            r = ws.max_row + 1
            for name_, value in values.items():
                if name_ in col:
                    ws.cell(row=r, column=col[name_], value=value)
            existing.add(key)
            next_no += 1
            added += 1
            print(f"added: {name[:60]}", file=sys.stderr)

    if added:
        wb.save(SHEET)
    print(json.dumps({"added": added}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
